"""Régressions hors ligne : aucune boîte, aucun token ni SQLite personnel."""

import base64
import contextlib
import csv
import io
import json
import tempfile
import unittest
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import MagicMock, patch
from openpyxl import load_workbook
import web
from app import database, importer, jobs, reclassifier
from app.connectors import gmail, microsoft, registry, status
from app.connectors.common import html_to_text
from app.exports import EXPORT_HEADERS
from app.models import DetectedEmail
from app.presentation import application_statistics, format_datetime
from app import settings

NOW = datetime(2026, 9, 16, 12, tzinfo=timezone.utc)


def fake_email(message_id="<fixture@example.invalid>", **changes):
    email = DetectedEmail(
        mailbox="Fixture",
        date=NOW,
        sender="Recrutement <rh@example.invalid>",
        subject="Votre candidature au poste de Développeur Python",
        message_id=message_id,
        score=10,
        status="RECEIVED",
        reasons=["fixture"],
        body="Nous avons bien reçu votre candidature.",
    )
    return replace(email, **changes)


class OfflineCase(unittest.TestCase):
    def setUp(self):
        self.stack = contextlib.ExitStack()
        self.addCleanup(self.stack.close)
        self.root = Path(self.stack.enter_context(tempfile.TemporaryDirectory()))
        self.stack.enter_context(
            patch.object(database, "DATABASE_PATH", self.root / "test.db")
        )
        for module, attribute, filename in (
            (gmail, "TOKEN_FILE", "gmail-token.json"),
            (gmail, "CREDENTIALS_FILE", "gmail-client.json"),
            (microsoft, "TOKEN_CACHE_FILE", "microsoft-cache.json"),
            (registry, "CREDENTIALS_FILE", "gmail-client.json"),
            (registry, "TOKEN_FILE", "gmail-token.json"),
            (status, "CREDENTIALS_FILE", "gmail-client.json"),
            (status, "TOKEN_FILE", "gmail-token.json"),
            (status, "TOKEN_CACHE_FILE", "microsoft-cache.json"),
        ):
            self.stack.enter_context(
                patch.object(
                    module,
                    attribute,
                    self.root / filename,
                )
            )
            self.stack.enter_context(
                patch.object(module, attribute, self.root / filename)
            )
        self.stack.enter_context(
            patch(
                "socket.socket.connect",
                side_effect=AssertionError("Réseau interdit dans les tests"),
            )
        )
        self.stack.enter_context(patch.dict(web.app.config, TESTING=True))
        self.stack.enter_context(contextlib.redirect_stdout(io.StringIO()))
        database.init_database()
        self.client = web.app.test_client()

    def application(
        self,
        company: str = "Entreprise Fictive",
        job_title: str = "Développeur Python",
        source: str = "Fixture",
        status: str = "RECEIVED",
        date: datetime = NOW,
    ) -> int:
        return database.create_application(
            company=company,
            job_title=job_title,
            source=source,
            status=status,
            date=date,
        )


class StorageAndImportTests(OfflineCase):
    def test_schema_upgrade_is_idempotent(self):
        with database.get_connection() as connection:
            connection.execute("DROP TABLE emails")
            connection.execute("DROP TABLE applications")
            connection.execute(
                "CREATE TABLE applications (id INTEGER PRIMARY KEY, company TEXT NOT NULL, job_title TEXT, source TEXT, first_seen TEXT NOT NULL, last_update TEXT NOT NULL, current_status TEXT NOT NULL)"
            )
            connection.execute(
                "CREATE TABLE emails (id INTEGER PRIMARY KEY, message_id TEXT UNIQUE NOT NULL, application_id INTEGER, mailbox TEXT NOT NULL, sender TEXT, subject TEXT, received_at TEXT NOT NULL, detected_status TEXT NOT NULL, score INTEGER NOT NULL)"
            )
        application_id = self.application()
        database.init_database()
        database.init_database()
        self.assertIsNotNone(database.get_application(application_id))
        self.assertTrue(database.save_email(fake_email(), application_id))
        self.assertEqual(
            database.get_application_emails(application_id)[0]["body"],
            fake_email().body,
        )

        database.delete_application(application_id)
        with database.get_connection() as connection:
            self.assertIsNone(
                connection.execute("SELECT application_id FROM emails").fetchone()[0]
            )
            import sqlite3

            with self.assertRaises(sqlite3.IntegrityError):
                connection.execute("UPDATE emails SET application_id = 123456")

    def test_import_deduplicates_sources_and_repeated_runs(self):

        email = fake_email()

        with (
            patch.object(
                importer,
                "scan_external_connectors",
                return_value=[
                    email,
                    replace(
                        email,
                        mailbox="Gmail API",
                    ),
                    fake_email(""),
                ],
            ),
            patch.object(
                importer,
                "extract_application_data",
                return_value=(
                    "Entreprise Fictive",
                    "Développeur Python",
                    "Fixture",
                ),
            ),
        ):
            self.assertEqual(
                importer.import_emails(),
                importer.ImportResult(1, 1, 1),
            )

            self.assertEqual(
                importer.import_emails(),
                importer.ImportResult(1, 0, 0),
            )

    def test_merge_and_delete_keep_emails(self):
        source = self.application("Source Fictive")
        target = self.application("Cible Fictive")
        database.save_email(fake_email(), source)
        database.merge_applications(source, target)
        self.assertIsNone(database.get_application(source))
        self.assertEqual(len(database.get_application_emails(target)), 1)
        database.delete_application(target)
        self.assertTrue(database.email_exists(fake_email().message_id))
        self.assertIsNone(database.get_application(target))

    def test_manual_status_survives_reclassification(self):
        application_id = self.application()
        database.save_email(fake_email(status="REJECTED"), application_id)
        database.set_manual_status(application_id, "INTERVIEW", "Note fictive")
        self.assertEqual(reclassifier.recompute_application_statuses(), 1)
        row = database.get_application(application_id)
        assert row is not None
        self.assertEqual(row["current_status"], "REJECTED")
        self.assertEqual(row["manual_status"], "INTERVIEW")
        self.assertTrue(row["manual_override"])

class WebTests(OfflineCase):
    def test_dashboard_displays_each_application_once(self):
        application_id = self.application()
        html = self.client.get("/").get_data(as_text=True)
        self.assertEqual(html.count(f'href="/application/{application_id}"'), 1)
        self.assertEqual(html.count("<!DOCTYPE html>"), 1)

    def test_search_status_and_exports_share_rows(self):
        selected = self.application("École Fictive")
        self.application("Autre Entreprise")
        database.set_manual_status(selected, "INTERVIEW", "Rappel UNIQUE")
        query = "?q=unique&status=INTERVIEW"
        html = self.client.get("/" + query).get_data(as_text=True)
        self.assertIn("École Fictive", html)
        self.assertNotIn("Autre Entreprise", html)
        csv_response = self.client.get("/export/csv" + query)
        csv_rows = list(
            csv.reader(
                io.StringIO(csv_response.data.decode("utf-8-sig")), delimiter=";"
            )
        )
        self.assertEqual(csv_rows[0], list(EXPORT_HEADERS))
        self.assertEqual(len(csv_rows), 2)
        excel_response = self.client.get("/export/xlsx" + query)
        workbook = load_workbook(io.BytesIO(excel_response.data))
        try:
            sheet = workbook.active
            assert sheet is not None
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            self.assertEqual(rows, csv_rows)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet["A1"].font.bold)
        finally:
            workbook.close()

    def test_detail_and_create_edit_delete_routes(self):
        response = self.client.post(
            "/application/new",
            data={
                "company": "Entreprise Fictive",
                "job_title": "Développeur Python",
                "status": "SENT",
                "note": "Note fictive",
            },
        )
        self.assertEqual(response.status_code, 302)
        application_id = int(response.headers["Location"].rsplit("/", 1)[1])
        detail = self.client.get(response.headers["Location"])
        self.assertEqual(detail.status_code, 200)
        self.assertIn(b"application.js", detail.data)
        response = self.client.post(
            f"/application/{application_id}/edit",
            data={
                "company": "Entreprise Modifiée",
                "job_title": "Développeur Python",
                "source": "Fixture",
                "status": "INTERVIEW",
                "note": "Nouvelle note",
            },
        )
        self.assertEqual(response.status_code, 302)
        application = database.get_application(application_id)
        if application is None:
            self.fail("Application was not found after creation")
        self.assertEqual(
            application["manual_status"], "INTERVIEW"
        )
        self.assertEqual(
            self.client.post(f"/application/{application_id}/delete").status_code, 302
        )
        self.assertEqual(
            self.client.get(f"/application/{application_id}").status_code, 404
        )

    def test_invalid_forms_and_missing_application(self):
        self.assertEqual(
            self.client.post(
                "/application/new", data={"company": "", "status": "SENT"}
            ).status_code,
            400,
        )
        self.assertEqual(
            self.client.post(
                "/application/new", data={"company": "Fixture", "status": "INVALID"}
            ).status_code,
            400,
        )
        self.assertEqual(self.client.get("/application/9999").status_code, 404)
        self.assertEqual(self.client.get("/application/new").status_code, 200)

    def test_connector_routes_keep_feedback_and_actions(self):
        with patch.object(gmail, "scan_gmail", return_value=[fake_email()]) as scan:
            response = self.client.post("/connectors/gmail/test", follow_redirects=True)
            self.assertEqual(response.status_code, 200)
            self.assertIn("En attente", response.get_data(as_text=True))
            scan.assert_not_called()
            self.assertTrue(jobs.run_next())
            scan.assert_called_once_with (settings.API_START_DATE)
        microsoft.TOKEN_CACHE_FILE.write_text("{}")
        with patch.object(microsoft, "scan_microsoft", return_value=[]):
            self.assertEqual(
                self.client.post("/connectors/microsoft/reconnect").status_code, 303
            )
            self.assertTrue(microsoft.TOKEN_CACHE_FILE.exists())
            jobs.run_next()
        self.assertFalse(microsoft.TOKEN_CACHE_FILE.exists())

    def test_scan_and_reclassify_routes(self):
        with patch.object(
            importer, "import_emails", return_value=importer.ImportResult(1, 1, 1)
        ) as scan:
            self.assertEqual(self.client.post("/scan").status_code, 303)
            self.client.post("/scan")
            scan.assert_not_called()
            self.assertEqual(len(jobs.list_jobs()), 1)
            jobs.run_next()
            scan.assert_called_once()
        result = reclassifier.ReclassifyResult(0, 0, 0, 0, 0)
        with patch.object(
            reclassifier, "reclassify_emails", return_value=result
        ) as reclassify:
            self.assertEqual(self.client.post("/reclassify").status_code, 303)
            reclassify.assert_not_called()
            jobs.run_next()
            reclassify.assert_called_once()

    def test_statistics_and_date_fallback(self):
        application_id = self.application()

        database.set_manual_status(
            application_id,
            "INTERVIEW",
            "",
        )

        stats = application_statistics(
            database.get_applications()
        )

        self.assertEqual(
            (
                stats["total"],
                stats["INTERVIEW"],
            ),
            (
                1,
                1,
            ),
        )

        self.assertEqual(
            format_datetime(None),
            "-",
        )

        self.assertEqual(
            format_datetime("date invalide"),
            "date invalide",
    )

class ConnectorTests(OfflineCase):
    def test_common_html_and_message_wrappers(self):
        html = "<style>hidden</style><p>Candidature &amp; entretien</p><script>hidden</script><br>Merci"
        expected = "Candidature & entretien\n\nMerci"
        self.assertEqual(html_to_text(html), expected)
        self.assertEqual(gmail.html_to_text(html), expected)
        self.assertEqual(microsoft.html_to_text(html), expected)
        message = gmail.build_email_message(
            "rh@example.invalid",
            "Sujet",
            "Wed, 16 Sep 2026 12:00:00 +0000",
            "<fixture@example.invalid>",
            "Corps",
        )
        self.assertIsNotNone(message["Date"])
        message = microsoft.build_email_message(
            "rh@example.invalid", "Sujet", "<fixture@example.invalid>", "Corps"
        )
        self.assertIsNone(message["Date"])
        self.assertEqual(message.get_content().strip(), "Corps")

    def test_gmail_pagination_and_fallback_message_id(self):
        service = MagicMock()
        messages = service.users.return_value.messages.return_value
        messages.list.return_value.execute.side_effect = [
            {"messages": [{"id": "first"}], "nextPageToken": "next"},
            {"messages": [{"id": "second"}, {}]},
        ]
        encoded = base64.urlsafe_b64encode(b"Votre candidature").decode()
        messages.get.return_value.execute.return_value = {
            "payload": {
                "mimeType": "text/plain",
                "body": {"data": encoded},
                "headers": [
                    {"name": "Subject", "value": "Votre candidature"},
                    {"name": "Date", "value": "Wed, 16 Sep 2026 12:00:00 +0000"},
                ],
            }
        }
        with (
            patch.object(gmail, "get_credentials", return_value=object()),
            patch.object(gmail, "build", return_value=service),
            patch.object(
                gmail, "score_message", return_value=(10, [], "Votre candidature")
            ),
        ):
            emails = gmail.scan_gmail(NOW)
        self.assertEqual(
            [email.message_id for email in emails], ["gmail:first", "gmail:second"]
        )
        self.assertEqual(messages.list.call_args_list[1].kwargs["pageToken"], "next")

    def test_microsoft_pagination_preserves_next_link(self):
        responses = []
        for identifier in ["first", "second"]:
            response = MagicMock()
            data: dict[str, object] = {
                "value": [
                    {
                        "id": identifier,
                        "subject": "Votre candidature",
                        "receivedDateTime": "2026-09-16T12:00:00Z",
                        "body": {
                            "contentType": "html",
                            "content": "<p>Votre candidature</p>",
                        },
                    }
                ]
            }
            if identifier == "first":
                data["@odata.nextLink"] = (
                    "https://graph.microsoft.com/v1.0/me/messages?$skiptoken=fixture"
                )
            response.json.return_value = data
            responses.append(response)
        with (
            patch.object(
                microsoft, "get_access_token", return_value="fictitious-token"
            ),
            patch.object(microsoft.requests, "get", side_effect=responses) as get,
            patch.object(
                microsoft, "score_message", return_value=(10, [], "Votre candidature")
            ),
        ):
            emails = microsoft.scan_microsoft(NOW)
        self.assertEqual(
            [email.message_id for email in emails],
            ["microsoft:first", "microsoft:second"],
        )
        self.assertIsNone(get.call_args_list[1].kwargs["params"])

    def test_registry_isolates_failures(self):
        registry.CREDENTIALS_FILE.write_text("{}")
        with (
            patch.object(registry, "MICROSOFT_CLIENT_ID", "fixture"),
            patch.object(registry, "scan_gmail", side_effect=RuntimeError("fixture")),
            patch.object(registry, "scan_microsoft", return_value=[fake_email()]),
        ):
            self.assertEqual(registry.scan_external_connectors(NOW), [fake_email()])

    def test_registry_skips_unconfigured_sources(self):
        with (
            patch.object(registry, "MICROSOFT_CLIENT_ID", ""),
            patch.object(registry, "scan_gmail") as gmail_scan,
            patch.object(registry, "scan_microsoft") as microsoft_scan,
        ):
            self.assertEqual(registry.scan_external_connectors(NOW), [])
            gmail_scan.assert_not_called()
            microsoft_scan.assert_not_called()

    def test_token_cache_round_trip_uses_temporary_file(self):
        microsoft.TOKEN_CACHE_FILE.write_text(
            json.dumps({"Account": {"fixture": {"username": "test@example.invalid"}}})
        )
        cache = microsoft.load_cache()
        self.assertIn("fixture", json.loads(cache.serialize())["Account"])
        cache.has_state_changed = True
        microsoft.save_cache(cache)
        self.assertIn(
            "fixture", json.loads(microsoft.TOKEN_CACHE_FILE.read_text())["Account"]
        )


if __name__ == "__main__":
    unittest.main()


class EvolutionTests(OfflineCase):
    def test_foreign_keys_and_orphan_repair(self):
        import sqlite3

        app_id = self.application()
        database.save_email(fake_email(), app_id)
        database.delete_application(app_id)
        with database.get_connection() as connection:
            self.assertEqual(connection.execute("PRAGMA foreign_keys").fetchone()[0], 1)
            self.assertIsNone(
                connection.execute("SELECT application_id FROM emails").fetchone()[0]
            )
            self.assertEqual(
                connection.execute("SELECT count(*) FROM emails").fetchone()[0], 1
            )
            with self.assertRaises(sqlite3.IntegrityError):
                connection.execute("UPDATE emails SET application_id = 999")
        # Simule une ancienne écriture effectuée sans intégrité activée.
        with sqlite3.connect(database.DATABASE_PATH) as connection:
            connection.execute("DROP TRIGGER emails_application_update")
            connection.execute("UPDATE emails SET application_id = 999")
        database.init_database()
        database.init_database()
        with database.get_connection() as connection:
            self.assertIsNone(
                connection.execute("SELECT application_id FROM emails").fetchone()[0]
            )
            self.assertEqual(
                connection.execute("PRAGMA foreign_key_check").fetchall(), []
            )

    def test_reanalysis_all_sources_preserves_manual_status(self):
        app_id = self.application()
        database.set_manual_status(app_id, "INTERVIEW", "Note privée fictive")
        for index, source in enumerate(["Gmail API", "Microsoft Graph", "IMAP"]):
            database.save_email(
                fake_email(f"<{index}@example.invalid>", mailbox=source), app_id
            )
        with (
            patch.object(reclassifier, "detect_status", return_value="REJECTED"),
        ):
            result = reclassifier.reclassify_emails()
        self.assertEqual((result.found, result.changed, result.missing), (3, 3, 0))
        row = database.get_application(app_id)
        assert row is not None
        self.assertEqual(row["manual_status"], "INTERVIEW")
        self.assertEqual(row["manual_note"], "Note privée fictive")
        self.assertEqual(row["current_status"], "REJECTED")

    def test_missing_body_is_counted_as_missing(self):
        app_id = self.application()

        database.save_email(
            fake_email(
                "<missing@fixture>",
                body="",
            ),
            app_id,
        )

        database.save_email(
            fake_email(
                "<available@fixture>",
                body="Nous avons bien reçu votre candidature.",
            ),
            app_id,
        )

        result = reclassifier.reclassify_emails()

        self.assertEqual(
            (
                result.scanned,
                result.found,
                result.missing,
            ),
            (
                2,
                1,
                1,
            ),
        )

    def test_queue_failure_and_serialization(self):
        jobs.init_jobs()
        first = jobs.enqueue("scan")
        self.assertEqual(jobs.enqueue("scan"), first)
        jobs.enqueue("reclassify")
        with database.get_connection() as connection:
            connection.execute(
                "UPDATE jobs SET state = 'running' WHERE id = ?", (first,)
            )
        with patch.object(jobs, "execute") as execute:
            self.assertFalse(jobs.run_next())
            execute.assert_not_called()
        with database.get_connection() as connection:
            connection.execute(
                "UPDATE jobs SET state = 'queued' WHERE id = ?", (first,)
            )
        with patch.object(jobs, "execute", side_effect=RuntimeError("secret-fixture")):
            self.assertTrue(jobs.run_next())
        with patch.object(jobs, "execute", return_value={"found": 1}):
            self.assertTrue(jobs.run_next())
        rows = jobs.list_jobs()
        self.assertEqual({row["state"] for row in rows}, {"failed", "succeeded"})
        self.assertNotIn(
            "secret-fixture", self.client.get("/jobs").get_data(as_text=True)
        )
        jobs.init_jobs()
        self.assertEqual(len(jobs.list_jobs()), 2)

    def test_settings_paths_and_dates(self):
        import os

        from app import settings

        with patch.dict(os.environ, {"GMAIL_TOKEN_FILE": "tokens/fixture.json"}):
            self.assertEqual(
                settings.configured_path("GMAIL_TOKEN_FILE", ""),
                settings.PROJECT_ROOT / "tokens/fixture.json",
            )
        configured = settings.configured_date(
            "2026-08-01T00:00:00+00:00"
        )

        offset = configured.utcoffset()

        assert offset is not None

        self.assertEqual(
            offset.total_seconds(),
            0,
        )
        self.assertIsNotNone(settings.configured_date("2026-08-01").tzinfo)
        with self.assertRaises(ValueError):
            settings.configured_date("incorrect")

    def test_worker_restart_keeps_pending_and_marks_interrupted(self):
        jobs.init_jobs()
        interrupted = jobs.enqueue("scan")
        jobs.enqueue("reclassify")
        with database.get_connection() as connection:
            connection.execute(
                "UPDATE jobs SET state = 'running' WHERE id = ?", (interrupted,)
            )
        jobs.recover_interrupted()
        self.assertEqual(
            {row["state"] for row in jobs.list_jobs()}, {"failed", "queued"}
        )
        with patch.object(jobs, "execute", return_value={"found": 0}):
            self.assertTrue(jobs.run_next())
        self.assertEqual(
            {row["state"] for row in jobs.list_jobs()}, {"failed", "succeeded"}
        )

    def test_http_stays_available_during_worker(self):
        import threading
        from concurrent.futures import ThreadPoolExecutor

        entered = threading.Event()
        release = threading.Event()

        def slow_task(kind):
            entered.set()
            if not release.wait(5):
                raise TimeoutError("Fixture bloquée")
            return {"detected": 0}

        self.client.post("/scan")
        with (
            patch.object(jobs, "execute", side_effect=slow_task),
            ThreadPoolExecutor(1) as pool,
        ):
            future = pool.submit(jobs.run_next)
            try:
                self.assertTrue(entered.wait(2))
                self.assertEqual(self.client.get("/").status_code, 200)
                self.assertEqual(self.client.get("/jobs").status_code, 200)
                self.assertFalse(future.done())
                self.assertFalse(jobs.run_next())
            finally:
                release.set()
            self.assertTrue(future.result(timeout=2))

    def test_dotenv_loaded_outside_project_and_environment_wins(self):
        import os
        import subprocess
        import sys

        from app import settings

        module_dir = self.root / "app"
        module_dir.mkdir()
        config = module_dir / "settings.py"
        config.write_text(Path(settings.__file__).read_text())
        (self.root / ".env").write_text(
            "MICROSOFT_CLIENT_ID=fixture-file\nGMAIL_TOKEN_FILE=tokens/fixture.json\nSCAN_START_DATE=2026-07-01\n"
        )
        script = "import runpy, sys; s=runpy.run_path(sys.argv[1]); print(s['MICROSOFT_CLIENT_ID']); print(s['GMAIL_TOKEN_FILE']); print(s['API_START_DATE'].isoformat())"
        environment = {
            key: value
            for key, value in os.environ.items()
            if key not in {"MICROSOFT_CLIENT_ID", "GMAIL_TOKEN_FILE", "SCAN_START_DATE"}
        }
        result = subprocess.check_output(
            [sys.executable, "-c", script, str(config)],
            cwd="/tmp",
            env=environment,
            text=True,
        )
        self.assertIn("fixture-file", result)
        self.assertIn(str(self.root / "tokens/fixture.json"), result)
        self.assertIn("2026-07-01T00:00:00+02:00", result)
        environment["MICROSOFT_CLIENT_ID"] = "fixture-process"
        result = subprocess.check_output(
            [sys.executable, "-c", script, str(config)],
            cwd="/tmp",
            env=environment,
            text=True,
        )
        self.assertIn("fixture-process", result)
