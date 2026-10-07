import csv
import io

from openpyxl import load_workbook
from unittest.mock import patch

from app import (
    database,
    importer,
    jobs,
    reclassifier,
)
from app.exports import (
    EXPORT_HEADERS,
)
from app.presentation import (
    application_statistics,
    format_datetime,
)

from tests.helpers import (
    OfflineCase,
)

class WebTests(OfflineCase):
    def test_dashboard_displays_each_application_once(self):
        application_id = self.application()
        html = self.client.get("/").get_data(as_text=True)
        self.assertEqual(html.count(f'href="/application/{application_id}"'), 1)
        self.assertEqual(html.count("<!DOCTYPE html>"), 1)

    def test_search_status_and_exports_share_rows(self):
        selected = self.application("École Fictive")
        self.application("Autre Entreprise")
        database.set_manual_status(selected, "INTERVIEW", "Rappel UNIQUE")
        query = "?q=unique&status=INTERVIEW"
        html = self.client.get("/" + query).get_data(as_text=True)
        self.assertIn("École Fictive", html)
        self.assertNotIn("Autre Entreprise", html)
        csv_response = self.client.get("/export/csv" + query)
        csv_rows = list(
            csv.reader(
                io.StringIO(csv_response.data.decode("utf-8-sig")), delimiter=";"
            )
        )
        self.assertEqual(csv_rows[0], list(EXPORT_HEADERS))
        self.assertEqual(len(csv_rows), 2)
        excel_response = self.client.get("/export/xlsx" + query)
        workbook = load_workbook(io.BytesIO(excel_response.data))
        try:
            sheet = workbook.active
            assert sheet is not None
            rows = [
                [
                    ""
                    if value is None
                    else value
                    for value in row
                ]
                for row in sheet.iter_rows(
                    values_only=True
                )
            ]
            self.assertEqual(rows, csv_rows)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet["A1"].font.bold)
        finally:
            workbook.close()

    def test_detail_and_create_edit_delete_routes(self):
        response = self.client.post(
            "/application/new",
            data={
                "company": "Entreprise Fictive",
                "job_title": "Développeur Python",
                "status": "SENT",
                "note": "Note fictive",
            },
        )
        self.assertEqual(response.status_code, 302)
        application_id = int(response.headers["Location"].rsplit("/", 1)[1])
        detail = self.client.get(response.headers["Location"])
        self.assertEqual(detail.status_code, 200)
        self.assertIn(b"application.js", detail.data)
        response = self.client.post(
            f"/application/{application_id}/edit",
            data={
                "company": "Entreprise Modifiée",
                "job_title": "Développeur Python",
                "source": "Fixture",
                "status": "INTERVIEW",
                "note": "Nouvelle note",
            },
        )
        self.assertEqual(response.status_code, 302)
        application = database.get_application(application_id)
        if application is None:
            self.fail("Application was not found after creation")
        self.assertEqual(
            application["manual_status"], "INTERVIEW"
        )
        self.assertEqual(
            self.client.post(f"/application/{application_id}/delete").status_code, 302
        )
        self.assertEqual(
            self.client.get(f"/application/{application_id}").status_code, 404
        )

    def test_invalid_forms_and_missing_application(self):
        self.assertEqual(
            self.client.post(
                "/application/new", data={"company": "", "status": "SENT"}
            ).status_code,
            400,
        )
        self.assertEqual(
            self.client.post(
                "/application/new", data={"company": "Fixture", "status": "INVALID"}
            ).status_code,
            400,
        )
        self.assertEqual(self.client.get("/application/9999").status_code, 404)
        self.assertEqual(self.client.get("/application/new").status_code, 200)

    def test_imap_test_route_enqueues_job(self):
        account_id = database.create_imap_account(
            email="test@example.invalid",
            provider="imap",
            host="imap.example.invalid",
            port=993,
            use_ssl=True,
            username="test@example.invalid",
            folder="INBOX",
            auth_method="password",
        )

        response = self.client.post(
            f"/connectors/imap/{account_id}/test",
            follow_redirects=True,
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        jobs_list = jobs.list_jobs()

        self.assertEqual(
            len(jobs_list),
            1,
        )

        self.assertEqual(
            jobs_list[0]["kind"],
            f"test:imap:{account_id}",
        )

    def test_scan_and_reclassify_routes(self):
        with patch.object(
            importer, "import_emails", return_value=importer.ImportResult(1, 1, 1)
        ) as scan:
            self.assertEqual(self.client.post("/scan").status_code, 303)
            self.client.post("/scan")
            scan.assert_not_called()
            self.assertEqual(len(jobs.list_jobs()), 1)
            jobs.run_next()
            scan.assert_called_once()
        result = reclassifier.ReclassifyResult(0, 0, 0, 0, 0)
        with patch.object(
            reclassifier, "reclassify_emails", return_value=result
        ) as reclassify:
            self.assertEqual(self.client.post("/reclassify").status_code, 303)
            reclassify.assert_not_called()
            jobs.run_next()
            reclassify.assert_called_once()

    def test_statistics_and_date_fallback(self):
        application_id = self.application()

        database.set_manual_status(
            application_id,
            "INTERVIEW",
            "",
        )

        stats = application_statistics(
            database.get_applications()
        )

        self.assertEqual(
            (
                stats["total"],
                stats["INTERVIEW"],
            ),
            (
                1,
                1,
            ),
        )

        self.assertEqual(
            format_datetime(None),
            "-",
        )

        self.assertEqual(
            format_datetime("date invalide"),
            "date invalide",
    )